"""
Выгрузка 10-значных кодов с путём по дереву — для сверки (label_dev) и других потребителей.

Зачем: разовая выгрузка «код;наименование» от 05.08 теряла путь, и в списке
выбора человек видел три раза «прочая». Здесь к каждому коду добавлены
наименования 6- и 8-значных родителей, полный путь и отметка, есть ли код
в действующем тарифе.

Источники:
  data/tnved.db             — дерево 2017 года + тариф TWS.BY (собирает parse_tnved.py);
  data/raw/tws_tnved.xlsx   — тариф TWS.BY: полная цепочка наименований (🠺)
                              и дата актуальности данных.

data_source:
  hierarchy      — код есть только в дереве 2017 года, в тарифе его нет;
  hierarchy+tws  — есть и в дереве, и в тарифе;
  tws            — есть только в тарифе (новые коды, в дереве 2017 года их нет).

path — цепочка наименований из тарифа, если код в тарифе (path_source=tariff);
иначе — наименования 6-, 8- и 10-значного уровня из дерева 2017 года (tree2017).

Запуск:  python export_codes.py   →   exports/tnved10_paths_<дата тарифа>.csv
Формат как у выгрузки 05.08: UTF-8 без BOM, разделитель «;», строки CRLF.
"""

from __future__ import annotations

import csv
import re
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).parent
DB_PATH = BASE / "data" / "tnved.db"
TWS_PATH = BASE / "data" / "raw" / "tws_tnved.xlsx"
OUT_DIR = BASE / "exports"
SEP = " › "

COLUMNS = [
    "tnved10", "name", "parent6_name", "parent8_name", "path", "path_source",
    "data_source", "in_tariff", "duty_rate", "tariff_as_of", "db_built_at",
]


def clean(text: str | None) -> str:
    """Схлопывает повторные пробелы — в дереве 2017 года их много."""
    return " ".join((text or "").split())


def load_tariff() -> tuple[dict[str, str], str]:
    """Цепочки наименований из тарифа TWS.BY и дата актуальности данных."""
    import openpyxl

    wb = openpyxl.load_workbook(TWS_PATH, read_only=True, data_only=True)
    as_of = ""
    if "Система TWS" in wb.sheetnames:
        for row in wb["Система TWS"].iter_rows(values_only=True):
            if row and clean(str(row[0] or "")) == "Актуальность данных":
                as_of = clean(str(row[1] or ""))

    chains: dict[str, str] = {}
    for row in wb["ТНВЭД"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = re.sub(r"\D", "", str(row[0]))
        if len(code) != 10:
            continue
        segments: list[str] = []
        for seg in str(row[1] or "").split("🠺"):
            seg = clean(seg).strip(":").strip()
            if seg and not (segments and seg.lower() == segments[-1].lower()):
                segments.append(seg)
        if segments:
            chains[code] = SEP.join(segments)
    wb.close()
    return chains, as_of


def main() -> None:
    for path in (DB_PATH, TWS_PATH):
        if not path.exists():
            sys.exit(f"Нет файла {path} — без него выгрузка была бы неполной, не делаю.")

    chains, tariff_as_of = load_tariff()
    with sqlite3.connect(DB_PATH) as conn:
        names = dict(conn.execute("SELECT code, description FROM codes"))
        built = dict(conn.execute("SELECT key, value FROM meta")).get("built_at", "")[:10]
        leaves = conn.execute(
            "SELECT code, description, duty_rate, data_source FROM codes "
            "WHERE level = 4 ORDER BY code"
        ).fetchall()

    # Отметка «в тарифе» берётся из базы; тариф в data/raw обязан быть тем же,
    # из которого собрана база, иначе отметки и пути разойдутся.
    in_db_tariff = {code for code, _, _, source in leaves if source != "hierarchy"}
    if in_db_tariff != set(chains):
        sys.exit(
            f"Тариф и база рассинхронизированы: в базе {len(in_db_tariff):,} кодов тарифа, "
            f"в {TWS_PATH.name} — {len(chains):,}. Пересоберите базу: python parse_tnved.py"
        )

    rows = []
    for code, name, duty, source in leaves:
        chain = chains.get(code)
        if chain:
            path, path_source = chain, "tariff"
        else:
            parts: list[str] = []
            for length in (6, 8, 10):
                desc = clean(names.get(code[:length]))
                if desc and not (parts and desc.lower() == parts[-1].lower()):
                    parts.append(desc)
            path, path_source = SEP.join(parts), "tree2017"
        rows.append([
            code,
            clean(name),
            clean(names.get(code[:6])),
            clean(names.get(code[:8])),
            path,
            path_source,
            source,
            0 if source == "hierarchy" else 1,
            duty or "",
            tariff_as_of,
            built,
        ])

    stamp = datetime.strptime(tariff_as_of, "%d.%m.%Y").strftime("%Y-%m-%d") if tariff_as_of else built
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f"tnved10_paths_{stamp}.csv"
    with open(out, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(COLUMNS)
        writer.writerows(rows)

    in_tariff = sum(r[7] for r in rows)
    print(f"{out}: {len(rows):,} кодов, в тарифе {in_tariff:,}, только в дереве 2017 года "
          f"{len(rows) - in_tariff:,}; тариф на {tariff_as_of}, база собрана {built}")


if __name__ == "__main__":
    main()
