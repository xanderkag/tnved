"""
Замер точности на размеченном наборе: описание товара → ожидаемый 10-значный код ТН ВЭД.

Строки уходят в развёрнутый сервис пакетом, как xlsx из UI (POST /api/classify/batch):
меряется то, что работает на сервере, — triage → classify без уточняющих вопросов.
В файле запроса только описание, ожидаемого кода в нём нет. Адрес сервиса — только
явно (--base) и только во внутренней сети: номенклатура наружу не уходит.

На 2/4/6/10 знаках считаются:
  - точность — код ответа совпал с ожидаемым в первых N знаках; строка без кода — промах;
  - базовая линия — всегда отвечать самым частым в наборе началом кода этой длины.
Ещё — ожидаемый код среди ответа и двух альтернатив («в тройке») и почему у строки нет
кода: модель назвала несуществующий, снятый или неполный код (А2 его не выдаёт), не
назвала кода, строка упала.

Наборы перекошены (у холдинга 78 % строк — один код), поэтому выборка по кодам: до
--per-code строк каждого ожидаемого кода. Итог по набору — с весами долей кодов в полном
наборе и 95 % интервалом; код, взятый целиком, разброса не даёт. --per-code 0 — весь
набор. Одинаковые описания уходят в сервис один раз.

Наборы (номенклатура холдинга, строки SLAI) в git не кладём: путь — аргументом,
результат — в eval_results/ (в .gitignore): построчный xlsx и сводка json.

  python eval.py "Датасет холдинга ТН ВЭД_ОКДП2.xlsx" --base http://10.10.13.10:8110
  python eval.py набор.xlsx --desc "Описание" --code "Код ТН ВЭД" --per-code 0 --base ...
  python eval.py набор.xlsx --dry-run    # состав выборки и базовая линия, без сервиса
"""

from __future__ import annotations

import argparse
import collections
import json
import math
import random
import re
import sys
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook

from classifier import _CODE_PROBLEM, DESIGNATION_ONLY
from netcheck import require_internal_url

LEVELS = (2, 4, 6, 10)
METRICS = [f"hit{n}" for n in LEVELS] + ["top3"]
BATCH_ROWS = 500  # не больше лимита сервиса на пакет (BATCH_MAX_ROWS)
HEADER_SCAN_ROWS = 10
OUT_DIR = Path(__file__).parent / "eval_results"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CODE_HEADER = re.compile(r"тн\s*вэд|\bhs\b|hs[ -]?code|commodity code", re.I)
DESC_HEADER = re.compile(r"описан|наимен|товар|descr|\bname\b|goods|product", re.I)

# Почему нет кода — по колонке «Ошибка» результата пакета; тексты — из classifier
NO_CODE_REASONS = {
    "несуществующий код": _CODE_PROBLEM["unknown"],
    "снятый код": _CODE_PROBLEM["retired"],
    "только обозначение": DESIGNATION_ONLY,
    "неполный код": _CODE_PROBLEM["not_leaf"],
    "модель не назвала код": "модель не назвала код",
}


def cell_text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def digits(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def code_of(value: object) -> str:
    """Ожидаемый код из ячейки: «8473 30 800 0» → «8473308000». Excel, хранящий код
    числом, теряет ведущий ноль у групп 01–09 — возвращаем."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    code = digits(value)
    return "0" + code if isinstance(value, int) and len(code) == 9 else code


def pct(x: float) -> str:
    return f"{100 * x:.1f}".replace(".", ",")


# ─── набор ────────────────────────────────────────────────────────────────────

def find_column(names: list[str], wanted: str | None, rx: re.Pattern, skip: int | None = None) -> int | None:
    for i, name in enumerate(names):
        if i != skip and name and (name.lower() == wanted.lower() if wanted else rx.search(name)):
            return i
    return None


def read_dataset(path: Path, sheet: str | None, desc: str | None, code: str | None) -> tuple[list[dict], dict]:
    """Строки с описанием и 10-значным кодом: [{row, description, expected}] + что отброшено."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    for h, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        names = [cell_text(c) for c in row or ()]
        ci = find_column(names, code, CODE_HEADER)
        di = find_column(names, desc, DESC_HEADER, skip=ci)
        if ci is not None and di is not None:
            break
    else:
        sys.exit("Не нашёл в первых строках колонки описания и кода — задайте --desc и --code")

    items, dropped = [], collections.Counter()
    for n, row in enumerate(rows[h + 1:], start=h + 2):
        row = row or ()
        description = cell_text(row[di]) if di < len(row) else ""
        expected = code_of(row[ci]) if ci < len(row) else ""
        if not description and not expected:
            continue
        if not expected:
            dropped["без кода"] += 1
        elif len(expected) != 10:
            dropped["код не 10 знаков"] += 1
        elif not description:
            dropped["без описания"] += 1
        else:
            items.append({"row": n, "description": description, "expected": expected})
    return items, {"header_row": h + 1, "description": names[di], "code": names[ci], "dropped": dict(dropped)}


def sample(items: list[dict], per_code: int, seed: int) -> list[dict]:
    """До per_code строк каждого ожидаемого кода, 0 — все."""
    by_code = collections.defaultdict(list)
    for it in items:
        by_code[it["expected"]].append(it)
    rnd = random.Random(seed)
    picked = []
    for code in sorted(by_code):
        rows = by_code[code]
        picked += rows if per_code <= 0 or len(rows) <= per_code else rnd.sample(rows, per_code)
    return picked


def baseline(items: list[dict]) -> dict:
    """Всегда отвечать самым частым в наборе началом кода длины n."""
    out = {}
    for n in LEVELS:
        prefix, count = collections.Counter(it["expected"][:n] for it in items).most_common(1)[0]
        out[f"hit{n}"] = {"prefix": prefix, "share": count / len(items)}
    return out


# ─── сервис ───────────────────────────────────────────────────────────────────

def classify_all(base: str, texts: list[str], poll: float) -> dict[str, dict]:
    """Описания → строка результата пакета (по заголовкам его xlsx). Пакеты по BATCH_ROWS строк."""
    results: dict[str, dict] = {}
    with httpx.Client(base_url=base, timeout=60) as client:
        for start in range(0, len(texts), BATCH_ROWS):
            chunk = texts[start:start + BATCH_ROWS]
            wb = Workbook()
            ws = wb.active
            ws.append(["Описание"])
            for i, t in enumerate(chunk, start=2):
                ws.cell(i, 1, t).data_type = "s"  # «=…» — текст, а не формула
            buf = BytesIO()
            wb.save(buf)
            r = client.post("/api/classify/batch", files={"file": ("eval.xlsx", buf.getvalue(), XLSX_TYPE)})
            if r.status_code != 200:
                sys.exit(f"Пакет не принят: {r.status_code} {r.text[:300]}")
            job, failures = r.json()["job_id"], 0
            while True:
                time.sleep(poll)
                try:
                    s = client.get(f"/api/classify/batch/{job}")
                    s.raise_for_status()
                    s, failures = s.json(), 0
                except httpx.HTTPError as e:  # пакет живёт на сервере — сбой опроса не повод бросать
                    failures += 1
                    if failures >= 10:
                        sys.exit(f"\nПакет {job}: сервис не отвечает на опрос ({e})")
                    continue
                print(f"\r  пакет {job}: {s['processed']}/{s['total']}, ошибок {s['errors_count']}",
                      end="", flush=True)
                if s["status"] in ("done", "failed"):
                    print()
                    break
            d = client.get(f"/api/classify/batch/{job}/download")
            d.raise_for_status()
            rows = list(load_workbook(BytesIO(d.content), read_only=True).active.iter_rows(values_only=True))
            for row in rows[1:]:
                res = dict(zip(rows[0], row))
                results[chunk[res["Строка файла"] - 2]] = res  # строка 1 — заголовок
    return results


# ─── оценка ───────────────────────────────────────────────────────────────────

def judge(item: dict, res: dict | None) -> dict:
    expected = item["expected"]
    res = res or {"Ошибка": "нет результата"}
    got = digits(res.get("Код ТН ВЭД"))
    alts = [digits(res.get("Альт. 1")), digits(res.get("Альт. 2"))]
    error = cell_text(res.get("Ошибка"))
    if got:
        status = "код выдан"
    elif error.startswith("код не выдан"):
        status = next((k for k, v in NO_CODE_REASONS.items() if v in error), "код не выдан")
    else:
        status = "ошибка сервиса"
    return {
        **item, "got": got, "alt1": alts[0], "alt2": alts[1],
        "confidence": cell_text(res.get("Уверенность")), "group": cell_text(res.get("Группа")),
        "status": status, "error": error,
        **{f"hit{n}": bool(got) and got[:n] == expected[:n] for n in LEVELS},
        "top3": expected in {got, *alts} - {""},
    }


def summarize(items: list[dict], judged: list[dict]) -> dict:
    """Точность по выборке и оценка на весь набор: доли кодов — веса, 95 % интервал."""
    total = collections.Counter(it["expected"] for it in items)
    n_all = len(items)
    by_code = collections.defaultdict(list)
    for j in judged:
        by_code[j["expected"]].append(j)

    estimate = {}
    for m in METRICS:
        est = var = 0.0
        for code, js in by_code.items():
            w, n, big_n = total[code] / n_all, len(js), total[code]
            p = sum(j[m] for j in js) / n
            est += w * p
            if n < big_n:  # код взят не целиком — разброс выборки с поправкой на конечный набор
                var += w * w * p * (1 - p) / max(n - 1, 1) * (big_n - n) / (big_n - 1)
        estimate[m] = {"value": est, "ci95": 1.96 * math.sqrt(var)}

    per_code = []
    for code, big_n in total.most_common():
        js = by_code.get(code, [])
        answers = collections.Counter(j["got"] or f"({j['status']})" for j in js).most_common(3)
        per_code.append({"code": code, "rows": big_n, "sampled": len(js),
                         **{m: sum(j[m] for j in js) for m in METRICS}, "answers": answers})
    return {
        "rows": n_all, "codes": len(total), "sampled": len(judged),
        "sample": {m: sum(j[m] for j in judged) / len(judged) for m in METRICS},
        "estimate": estimate,
        "baseline": baseline(items),
        "statuses": dict(collections.Counter(j["status"] for j in judged).most_common()),
        "per_code": per_code,
    }


def print_report(s: dict) -> None:
    head = "".join(f"{f'{n} зн.':>12}" for n in LEVELS) + f"{'в тройке':>12}"
    print(f"\n{'':26}{head}")
    print(f"{'по набору (веса кодов)':26}" + "".join(
        f"{pct(s['estimate'][m]['value']) + ('±' + pct(s['estimate'][m]['ci95']) if s['estimate'][m]['ci95'] else ''):>12}"
        for m in METRICS))
    print(f"{'по выборке':26}" + "".join(f"{pct(s['sample'][m]):>12}" for m in METRICS))
    print(f"{'всегда самое частое':26}" + "".join(
        f"{pct(s['baseline'][f'hit{n}']['share']):>12}" for n in LEVELS)
        + f"   ({', '.join(s['baseline'][f'hit{n}']['prefix'] for n in LEVELS)})")
    print(f"\nОтветы ({s['sampled']} строк выборки): "
          + ", ".join(f"{k} — {v}" for k, v in s["statuses"].items()))
    print("\nПо кодам: строк в наборе / в выборке, совпало на 4/6/10 знаках, в тройке; что ответил сервис")
    for c in s["per_code"]:
        answers = ", ".join(f"{a} ×{k}" for a, k in c["answers"])
        print(f"  {c['code']}  {c['rows']:>5}/{c['sampled']:<4} {c['hit4']:>3}/{c['hit6']}/{c['hit10']}, "
              f"{c['top3']:<4} {answers}")


def write_rows(path: Path, judged: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "eval"
    cols = ["row", "description", "expected", "got", "confidence", "alt1", "alt2", "group", "status",
            *METRICS, "error"]
    ws.append(["Строка набора", "Описание", "Ожидали", "Код ТН ВЭД", "Уверенность", "Альт. 1", "Альт. 2",
               "Группа", "Статус", *(f"Совпало, {n} зн." for n in LEVELS), "В тройке", "Ошибка"])
    for j in judged:
        ws.append([j.get(c) for c in cols])
        ws.cell(ws.max_row, 2).data_type = "s"
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("dataset", type=Path, help="xlsx: колонка описания и колонка ожидаемого кода")
    ap.add_argument("--base", help="адрес сервиса, например http://10.10.13.10:8110 — только внутренняя сеть")
    ap.add_argument("--sheet", help="лист (по умолчанию первый)")
    ap.add_argument("--desc", help="заголовок колонки описания (по умолчанию — узнаётся)")
    ap.add_argument("--code", help="заголовок колонки ожидаемого кода (по умолчанию — «ТН ВЭД», «HS code»)")
    ap.add_argument("--per-code", type=int, default=30, help="строк на ожидаемый код, 0 — все (30)")
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--poll", type=float, default=3.0, help="опрос пакета, с")
    ap.add_argument("--dry-run", action="store_true", help="только выборка и базовая линия, без сервиса")
    args = ap.parse_args()

    items, info = read_dataset(args.dataset, args.sheet, args.desc, args.code)
    if not items:
        sys.exit("В наборе нет строк с описанием и 10-значным кодом")
    picked = sample(items, args.per_code, args.seed)
    texts = list(dict.fromkeys(it["description"] for it in picked))
    dropped = ", ".join(f"{k} {v}" for k, v in info["dropped"].items()) or "нет"
    print(f"Набор {args.dataset.name}: колонки «{info['description']}» и «{info['code']}» "
          f"(заголовок в строке {info['header_row']}); строк с кодом {len(items)}, отброшено: {dropped}; "
          f"кодов {len({it['expected'] for it in items})}")
    print(f"Выборка: {f'до {args.per_code} строк на код' if args.per_code > 0 else 'весь набор'} — "
          f"{len(picked)} строк, в сервис {len(texts)} описаний")
    if args.dry_run:
        for n, b in baseline(items).items():
            print(f"  всегда самое частое, {n[3:]} зн.: {b['prefix']} — {pct(b['share'])} %")
        return
    if not args.base:
        sys.exit("Адрес сервиса не задан (--base): по умолчанию набор никуда не отправляем")
    try:
        base = require_internal_url(args.base, "--base").rstrip("/")
    except RuntimeError as e:
        sys.exit(str(e))

    try:
        r = httpx.get(f"{base}/health", timeout=15)
        r.raise_for_status()
        health = r.json()
    except httpx.HTTPError as e:
        sys.exit(f"Сервис {base} не готов: {e}")
    print(f"Сервис {base}: {health.get('mode')}, модель {health.get('llm_model')}, векторов {health.get('vectors')}")
    t0 = time.time()
    results = classify_all(base, texts, args.poll)
    wall = time.time() - t0
    judged = [judge(it, results.get(it["description"])) for it in picked]
    summary = summarize(items, judged)
    print_report(summary)

    OUT_DIR.mkdir(exist_ok=True)
    stem = f"{args.dataset.stem}_{datetime.now():%Y%m%d-%H%M}"
    write_rows(OUT_DIR / f"{stem}.xlsx", judged)
    (OUT_DIR / f"{stem}.json").write_text(json.dumps({
        "dataset": args.dataset.name, "columns": info, "per_code_limit": args.per_code, "seed": args.seed,
        "base": base, "health": health, "descriptions_sent": len(texts),
        "seconds": round(wall), "seconds_per_description": round(wall / len(texts), 1),
        **summary,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    per_text = f"{wall / len(texts):.1f}".replace(".", ",")
    print(f"\n{wall:.0f} с, {per_text} с на описание; результат — {OUT_DIR / stem}.xlsx / .json")


if __name__ == "__main__":
    main()
