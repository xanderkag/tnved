"""
Парсит сырые данные ТН ВЭД (CSV / XML / Excel) и сохраняет в SQLite.

Источники:
  data/raw/tnved.csv|xml — иерархия (группы → позиции → ... → листья)
  data/raw/tws_tnved.xlsx — свежие листья + ставка пошлины (TWS.BY, обновл. ежедневно)

Алгоритм:
  1) Читаем иерархию (если есть) → получаем все уровни кодов с описаниями.
  2) Читаем TWS.BY (если есть) → 10-значные листья с описанием и duty_rate.
  3) Сливаем: TWS.BY перетирает описание для совпадающих кодов и ставит duty_rate.
     Коды только в TWS.BY — добавляются как новые листья (level=4).
     Коды только в иерархии — остаются (без duty_rate).

Запуск: python parse_tnved.py
Результат: data/tnved.db
  - таблица codes(code, description, level, parent_code, full_path, duty_rate, data_source)
  - таблица meta(key, value) — хранит дату обновления / общую статистику
"""

from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
import sys
from pathlib import Path

RAW_DIR = Path(__file__).parent / "data" / "raw"
DB_PATH = Path(__file__).parent / "data" / "tnved.db"
TWS_PATH = RAW_DIR / "tws_tnved.xlsx"


# ─── helpers ──────────────────────────────────────────────────────────────────

def code_to_level(code: str) -> int:
    """Определяет уровень иерархии по длине кода."""
    length = len(code.strip())
    if length <= 2:
        return 1   # группа
    if length == 4:
        return 2   # позиция
    if length <= 6:
        return 3   # субпозиция
    if length <= 8:
        return 3   # промежуточный уровень (8 знаков) — тоже субпозиция
    return 4       # подсубпозиция (10 знаков) — листовой узел


def parent_code(code: str, known) -> str | None:
    """Ближайший предок, который есть в справочнике: для 10 знаков — 8, 6, 4, 2.

    Раньше родителем 10-значного всегда считался 6-значный, и путь терял 8-значный
    уровень: у 7318157008 пропадало «болты с шестигранной головкой из
    коррозионностойкой стали» — а это и отличает его от соседних кодов.
    """
    code = code.strip()
    for length in (8, 6, 4, 2):
        if length < len(code) and code[:length] in known:
            return code[:length]
    return None


PATH_SEP = " → "
# Раздел в начале цепочки тарифа: «Недрагоценные металлы и изделия из них (гр. 72-83)»
SECTION_RE = re.compile(r"\(гр\. \d")


def _norm(text: str | None) -> str:
    return " ".join((text or "").lower().replace("ё", "е").split()).strip(":").strip()


def _same_name(segment: str, tree_desc: str | None) -> bool:
    """Сегмент цепочки тарифа — это узел дерева? Длинные наименования в дереве обрезаны."""
    a, b = _norm(segment), _norm(tree_desc)
    return bool(b) and (a == b or (len(b) > 40 and a.startswith(b)))


# Пометки тарифа в конце сегмента: «(с 01.01.2022)» и ссылка на сноску «части 5)».
# Сноска — только после слова: «(кроме указанных в субпозиции 9603 30)» — это текст.
NOTE_RE = re.compile(r"\s*\((?:с|по|до)\s\d{2}\.\d{2}\.\d{4}\)$|(?<=[^\d\s])\s+\d{1,2}\)$")


def _clean_segment(seg: str) -> str:
    seg, prev = " ".join(seg.split()), None
    while seg != prev:
        prev, seg = seg, NOTE_RE.sub("", seg).rstrip(":").rstrip()
    return seg


def tws_chain(text: str) -> list[str]:
    """«… 🠺 мебель для сидения вращающаяся …: (с 01.01.2022) 🠺 прочая» → сегменты без пометок и повторов."""
    out: list[str] = []
    for seg in text.split("🠺"):
        seg = _clean_segment(seg)
        if seg and not (out and _norm(seg) == _norm(out[-1])):
            out.append(seg)
    return out


def chain_below_heading(chain: list[str], heading: str | None, group: str | None) -> list[str]:
    """Часть цепочки тарифа ниже товарной позиции (4 знака).

    Цепочка в тарифе начинается с разного уровня: с раздела, группы, позиции, а у 65 %
    кодов — сразу ниже позиции. Срезаем всё до позиции включительно; если позиции
    в цепочке нет — только раздел и группу в начале.
    """
    hits = [i for i, seg in enumerate(chain) if _same_name(seg, heading)]
    if hits:
        return chain[hits[-1] + 1:]
    start = 0
    while start < len(chain) - 1 and (SECTION_RE.search(chain[start]) or _same_name(chain[start], group)):
        start += 1
    return chain[start:]


def build_full_paths(rows: list[dict]) -> dict[str, str]:
    """code → «группа → позиция → … → код».

    Узлы дерева идут через ближайшего предка, с 8-значным уровнем. У кодов тарифа
    ниже позиции — цепочка самого тарифа: в ней все уровни с тире, и она в действующей
    редакции. Наименования 6/8 из дерева 2017 года местами устарели (у 854449
    «на напряжение не более 80 В» при действующем «не более 1000 В»), их к кодам
    тарифа не подставляем.
    """
    by_code = {r["code"]: r for r in rows}
    segs: dict[str, list[str]] = {}

    def add(parts: list[str], seg: str) -> None:
        if seg and not (parts and _norm(seg) == _norm(parts[-1])):
            parts.append(seg)

    def get_segs(code: str) -> list[str]:
        if code in segs:
            return segs[code]
        row = by_code[code]
        chain = row.get("tws_chain")
        if chain:
            heading, group = code[:4], code[:2]
            base = heading if heading in by_code else group if group in by_code else None
            parts = list(get_segs(base)) if base else []
            below = chain_below_heading(
                chain,
                by_code[heading]["description"] if heading in by_code else None,
                by_code[group]["description"] if group in by_code else None,
            )
            for seg in below:
                add(parts, seg)
        else:
            parent = parent_code(code, by_code)
            parts = list(get_segs(parent)) if parent else []
            add(parts, row["description"])
        segs[code] = parts
        return parts

    return {code: PATH_SEP.join(get_segs(code)) for code in by_code}


# ─── parsers ──────────────────────────────────────────────────────────────────

CODE_COLS = {"code", "kod", "код", "tnved", "code_tnved", "id"}
DESC_COLS = {"description", "name", "наименование", "название", "desc", "descr", "simple_nam", "nam", "наим"}


def parse_csv(path: Path) -> list[dict]:
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        # Пропускаем строки до настоящего заголовка (с именами столбцов)
        lines = f.readlines()

    header_idx = 0
    for i, line in enumerate(lines[:20]):
        cols = [c.strip().strip('"').lower() for c in line.split(",")]
        if any(c in CODE_COLS for c in cols) and any(c in DESC_COLS for c in cols):
            header_idx = i
            break

    content = "".join(lines[header_idx:])
    delimiter = ";" if content.count(";") > content.count(",") else ","

    import io
    reader = csv.DictReader(io.StringIO(content), delimiter=delimiter)
    fieldnames = reader.fieldnames or []

    code_col = next((h for h in fieldnames if h.strip().lower() in CODE_COLS), None)
    desc_col = next((h for h in fieldnames if h.strip().lower() in DESC_COLS), None)

    if not code_col or not desc_col:
        if len(fieldnames) >= 2:
            code_col, desc_col = fieldnames[0], fieldnames[1]
        else:
            raise ValueError(f"Не удалось определить столбцы в {path}. Заголовки: {fieldnames}")

    for row in reader:
        code = str(row.get(code_col, "")).strip().lstrip("'\"")
        desc = str(row.get(desc_col, "")).strip()
        if code and desc and code.isdigit():
            rows.append({"code": code, "description": desc})

    return rows


def parse_xml(path: Path) -> list[dict]:
    from lxml import etree

    tree = etree.parse(str(path))
    root = tree.getroot()
    rows = []

    # Пробуем несколько возможных структур XML
    # Структура 1: <TNVED><RAZDEL><GRUPPA><POZICIYA kod="..." naim="...">
    for elem in root.iter():
        tag = elem.tag.split("}")[-1].upper() if "}" in elem.tag else elem.tag.upper()
        code = elem.get("kod") or elem.get("code") or elem.get("KOD") or elem.get("CODE")
        desc = elem.get("naim") or elem.get("name") or elem.get("NAIM") or elem.get("NAME") or elem.text
        if code and desc and str(code).strip().isdigit():
            rows.append({"code": str(code).strip(), "description": str(desc).strip()})

    return rows


def parse_excel(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    rows_raw = list(ws.iter_rows(values_only=True))

    if not rows_raw:
        raise ValueError("Excel-файл пустой")

    # Ищем строку с заголовками
    header_row = 0
    for i, row in enumerate(rows_raw[:10]):
        row_str = [str(c).lower() if c else "" for c in row]
        if any(k in " ".join(row_str) for k in ("код", "code", "наим", "name")):
            header_row = i
            break

    headers = [str(c).lower().strip() if c else "" for c in rows_raw[header_row]]
    code_idx = next((i for i, h in enumerate(headers) if any(k in h for k in ("код", "code", "tnved"))), 0)
    desc_idx = next((i for i, h in enumerate(headers) if any(k in h for k in ("наим", "name", "описан", "descr"))), 1)

    rows = []
    for row in rows_raw[header_row + 1:]:
        if len(row) <= max(code_idx, desc_idx):
            continue
        code = str(row[code_idx]).strip() if row[code_idx] is not None else ""
        desc = str(row[desc_idx]).strip() if row[desc_idx] is not None else ""
        code = code.lstrip("'\"").rstrip(".0")  # убираем лишние символы Excel
        if code and desc and code.isdigit():
            rows.append({"code": code, "description": desc})

    wb.close()
    return rows


def parse_tws(path: Path) -> dict[str, dict]:
    """
    Парсит xlsx с TWS.BY (лист «ТНВЭД», столбцы: Код | Наименование | Тариф | Подробности).
    Возвращает {code: {description, duty_rate, full_path_tws}}.

    Если формат файла поменяется (другие имена столбцов / другой лист) — падаем
    с понятным сообщением, а не молча собираем кривую базу.
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    if "ТНВЭД" not in wb.sheetnames:
        raise ValueError(
            f"TWS.BY xlsx: ожидался лист «ТНВЭД», нашли: {wb.sheetnames}. "
            "Скорее всего формат файла поменялся — поправь parse_tws."
        )
    ws = wb["ТНВЭД"]

    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None) or ()
    header_norm = [str(c).strip().lower() if c else "" for c in header]
    expected = ("код", "наименование", "тариф")
    for word in expected:
        if not any(word in h for h in header_norm):
            raise ValueError(
                f"TWS.BY xlsx: в заголовке листа «ТНВЭД» нет столбца «{word}». "
                f"Реальный заголовок: {header}"
            )

    out: dict[str, dict] = {}
    for row in rows_iter:
        if not row or len(row) < 3:
            continue
        code, name, tariff, *_ = row
        if code is None or name is None:
            continue
        code = str(code).strip().lstrip("'\"")
        if not code.isdigit():
            continue
        full_path_tws = str(name).strip()
        # Описание листового кода — последний сегмент пути.
        # В TWS-выгрузке сегменты разделены символом " 🠺 ".
        leaf_desc = full_path_tws.split(" 🠺 ")[-1].strip().rstrip(":").strip()
        duty = str(tariff).strip() if tariff is not None else None
        out[code] = {
            "description": leaf_desc or full_path_tws,
            "duty_rate": duty,
            "full_path_tws": full_path_tws,
            "chain": tws_chain(full_path_tws),
        }

    wb.close()

    # Sentinel-check: если получили совсем мало кодов или почти ни у кого нет
    # тарифа — это сигнал что формат файла поменялся.
    if len(out) < 1000:
        raise ValueError(
            f"TWS.BY xlsx: получили только {len(out)} кодов (ожидаем >10к). "
            "Скорее всего формат изменился."
        )
    with_duty = sum(1 for v in out.values() if v["duty_rate"])
    if with_duty < len(out) * 0.5:
        raise ValueError(
            f"TWS.BY xlsx: только {with_duty}/{len(out)} кодов с тарифом. "
            "Колонка «Тариф» съехала?"
        )
    return out


def load_hierarchy() -> list[dict]:
    """Иерархия. Возвращает [] если файл не найден."""
    candidates = [
        (RAW_DIR / "tnved.csv",   parse_csv),
        (RAW_DIR / "tnved.xml",   parse_xml),
        (RAW_DIR / "tnved.xlsx",  parse_excel),
    ]
    for path, parser in candidates:
        if path.exists() and path.stat().st_size > 1000:
            print(f"Парсим иерархию: {path}")
            return parser(path)
    return []


def load_tws() -> dict[str, dict]:
    """Свежие листья с TWS.BY. Возвращает {} если файл не найден."""
    if not TWS_PATH.exists() or TWS_PATH.stat().st_size < 1000:
        return {}
    print(f"Парсим TWS.BY: {TWS_PATH}")
    return parse_tws(TWS_PATH)


def tws_as_of() -> str | None:
    """«Актуальность данных» тарифа TWS.BY (лист «Система TWS»), например «06.05.2026»; нет — None."""
    if not TWS_PATH.exists():
        return None
    from openpyxl import load_workbook

    wb = load_workbook(str(TWS_PATH), read_only=True, data_only=True)
    try:
        if "Система TWS" not in wb.sheetnames:
            return None
        for row in wb["Система TWS"].iter_rows(values_only=True):
            if row and " ".join(str(row[0] or "").split()) == "Актуальность данных":
                return " ".join(str(row[1] or "").split()) or None
        return None
    finally:
        wb.close()


# ─── DB ───────────────────────────────────────────────────────────────────────

def merge_sources(hier: list[dict], tws: dict[str, dict]) -> list[dict]:
    """
    Сливает иерархию и TWS-листья.
    Для совпадающих кодов TWS-описание перетирает иерархическое (свежее).
    Коды только в TWS добавляются как новые листья.
    """
    by_code: dict[str, dict] = {}
    for r in hier:
        code = r["code"]
        by_code[code] = {
            "code": code,
            "description": r["description"],
            "duty_rate": None,
            "data_source": "hierarchy",
        }
    for code, t in tws.items():
        if code in by_code:
            by_code[code]["description"] = t["description"]
            by_code[code]["duty_rate"] = t["duty_rate"]
            by_code[code]["data_source"] = "hierarchy+tws"
        else:
            by_code[code] = {
                "code": code,
                "description": t["description"],
                "duty_rate": t["duty_rate"],
                "data_source": "tws",
            }
        by_code[code]["tws_chain"] = t.get("chain") or []
    return list(by_code.values())


def save_to_db(rows: list[dict], tariff_as_of: str | None = None):
    # Собираем рядом и подменяем в конце: упавшая сборка не оставляет сервис без базы.
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = DB_PATH.with_name(DB_PATH.name + ".tmp")
    if tmp_path.exists():
        tmp_path.unlink()

    conn = sqlite3.connect(tmp_path)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE codes (
            code        TEXT PRIMARY KEY,
            description TEXT NOT NULL,
            level       INTEGER NOT NULL,
            parent_code TEXT,
            full_path   TEXT,
            duty_rate   TEXT,
            data_source TEXT
        )
    """)
    cur.execute("CREATE INDEX idx_level ON codes(level)")

    cur.execute("CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT)")

    print("Строим полные пути ...")
    paths = build_full_paths(rows)
    known = {r["code"] for r in rows}

    records = [
        (
            r["code"],
            r["description"],
            code_to_level(r["code"]),
            parent_code(r["code"], known),
            paths.get(r["code"], r["description"]),
            r.get("duty_rate"),
            r.get("data_source", "hierarchy"),
        )
        for r in rows
    ]

    cur.executemany(
        "INSERT OR IGNORE INTO codes VALUES (?,?,?,?,?,?,?)",
        records,
    )

    from datetime import datetime
    cur.execute("INSERT INTO meta VALUES (?,?)", ("built_at", datetime.utcnow().isoformat()))
    cur.execute("INSERT INTO meta VALUES (?,?)", ("total", str(len(records))))
    if tariff_as_of:  # дата тарифа для /api/codes; без неё поле пустое, а не дата сборки
        cur.execute("INSERT INTO meta VALUES (?,?)", ("tariff_as_of", tariff_as_of))

    conn.commit()

    total = cur.execute("SELECT COUNT(*) FROM codes").fetchone()[0]
    leaves = cur.execute("SELECT COUNT(*) FROM codes WHERE level=4").fetchone()[0]
    with_rate = cur.execute("SELECT COUNT(*) FROM codes WHERE duty_rate IS NOT NULL").fetchone()[0]
    conn.close()
    os.replace(tmp_path, DB_PATH)

    print(f"Сохранено записей: {total:,}  (листовых: {leaves:,}, со ставкой: {with_rate:,})")


# ─── main ─────────────────────────────────────────────────────────────────────

def main():
    hier = load_hierarchy()
    tws = load_tws()

    if not hier and not tws:
        print("Не нашёл ни иерархию, ни TWS.BY. Запустите fetch_tnved.py.")
        sys.exit(1)

    print(f"Иерархия: {len(hier):,} строк, TWS.BY: {len(tws):,} кодов")
    rows = merge_sources(hier, tws)

    if len(rows) < 100:
        print("Слишком мало записей после слияния — что-то пошло не так.")
        sys.exit(1)

    save_to_db(rows, tariff_as_of=tws_as_of() if tws else None)
    print(f"\nГотово: {DB_PATH}")
    print("Следующий шаг: python build_index.py")


if __name__ == "__main__":
    main()
