"""
Строит FAISS-индекс по описаниям кодов ТН ВЭД.

Запуск: python build_index.py
Результат:
  data/tnved_vecs.npy        — промежуточные векторы (переиспользуются)
  data/tnved.faiss           — векторный индекс
  data/tnved_meta.json       — список {code, description, full_path, ...}
  data/tnved_index_info.json — каким бэкендом и в какой размерности собран

Векторизацией занимается embedder.py — bge-m3 на GPU-сервере через
`/v1/embeddings` (EMBEDDINGS_BASE_URL; быстро, ничего локально не нужно), либо,
только явно, локальная e5 на CPU (EMBEDDER_BACKEND=local, requirements-e5.txt).
Подробнее — docstring embedder.py. Индекс собирается здесь, заранее: при сборке
Docker-образа он не строится, образ берёт готовые файлы из data/.

Прогон разведён на два этапа, каждый в своём процессе:
  embed — считает векторы, пишет .npy
  index — читает .npy, собирает FAISS
Причина разделения: локальный бэкенд тянет torch, и torch с faiss в одном
процессе на некоторых сборках роняют интерпретатор. Плюс так этап embed
возобновляем: считается кусками, каждый кусок сразу на диск.
Без аргументов скрипт запускает оба этапа сам, по очереди.
"""

import hashlib
import json
import os
import sqlite3
import subprocess
import sys
import time
from pathlib import Path

BASE = Path(__file__).parent
DB_PATH = BASE / "data" / "tnved.db"
VECS_PATH = BASE / "data" / "tnved_vecs.npy"
PARTS_DIR = BASE / "data" / "vecs_parts"
FAISS_PATH = BASE / "data" / "tnved.faiss"
META_PATH = BASE / "data" / "tnved_meta.json"
INFO_PATH = BASE / "data" / "tnved_index_info.json"
VECS_INFO_PATH = BASE / "data" / "tnved_vecs_info.json"
TWS_PATH = BASE / "data" / "raw" / "tws_tnved.xlsx"

# Насколько подрезать описание каждого уровня иерархии при склейке текста.
LEVEL_TEXT_LIMIT = int(os.environ.get("EMBED_LEVEL_LIMIT", "140"))
CHUNK_SIZE = int(os.environ.get("EMBED_CHUNK", "1024"))
# Актуально только для локального бэкенда: на слабом CPU многопоточность
# torch плюс нехватка памяти давали access violation.
EMBED_THREADS = int(os.environ.get("EMBED_THREADS", "8"))


def _part_path(offset: int, count: int) -> Path:
    return PARTS_DIR / f"part_{offset:06d}_{count}.npy"


def _load_tws_chains() -> dict[str, str]:
    """Полные цепочки наименований из выгрузки TWS.BY.

    В файле колонка «Наименование» содержит всю цепочку через символ 🠺:
    «… 🠺 мебель для сидения вращающаяся с регулирующими высоту
    приспособлениями: (с 01.01.2022) 🠺 прочая».

    Загрузчик `parse_tnved.py` брал из неё только последний сегмент, поэтому
    в базе у 9401390000 осталось одно слово «прочая», а различающая
    формулировка (по сути «офисное кресло») терялась. Здесь берём цепочку
    целиком — она и есть тот текст, по которому вообще можно искать.
    """
    if not TWS_PATH.exists():
        print(f"Выгрузки TWS не нашёл ({TWS_PATH.name}) — беру описания только из базы.")
        return {}

    try:
        import openpyxl
    except ImportError:
        print("openpyxl не установлен — цепочки TWS не подхватываю.")
        return {}

    import re

    wb = openpyxl.load_workbook(TWS_PATH, read_only=True, data_only=True)
    sheet = "ТНВЭД" if "ТНВЭД" in wb.sheetnames else wb.sheetnames[-1]
    chains: dict[str, str] = {}
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = re.sub(r"\D", "", str(row[0]))
        name = str(row[1] or "").strip()
        if len(code) != 10 or not name:
            continue
        segments: list[str] = []
        for seg in re.split(r"🠺", name):
            seg = seg.strip().strip(":").strip()
            if not seg:
                continue
            if len(seg) > LEVEL_TEXT_LIMIT:
                seg = seg[:LEVEL_TEXT_LIMIT].rsplit(" ", 1)[0] + "…"
            if segments and seg.lower() == segments[-1].lower():
                continue
            segments.append(seg)
        if segments:
            chains[code] = " → ".join(segments)
    wb.close()
    print(f"Цепочек наименований из TWS: {len(chains):,}")
    return chains


def build_index_texts(rows: list[dict]) -> list[str]:
    """Собирает текст для векторизации как цепочку от группы до самого кода.

    Зачем не брать full_path как есть: он заполнен у всех записей, но у части
    кодов (пришедших без иерархии) содержит только собственное название листа.
    В итоге 659 разных кодов индексировались одним словом «прочие», а всего
    неуникальный текст был у 22,8 % записей — такие коды векторный поиск
    различить не может в принципе.

    Поэтому текст склеиваем из описаний всех предков по префиксам кода
    (2/4/6/8/10 знаков): «МЕБЕЛЬ ДЛЯ СИДЕНИЯ … → мебель обитая → прочая».
    """
    by_code = {r["code"]: r for r in rows}
    tws_chains = _load_tws_chains()

    # Предки могут быть выше level>=3, поэтому названия берём из всей таблицы.
    conn = sqlite3.connect(DB_PATH)
    all_desc = dict(conn.execute("SELECT code, description FROM codes").fetchall())
    conn.close()

    texts = []
    used_tws = 0
    for row in rows:
        code = row["code"]

        # Цепочка из TWS полнее и свежее (редакция ГС-2022), поэтому в приоритете.
        chain = tws_chains.get(code)
        if chain:
            texts.append(chain)
            used_tws += 1
            continue

        parts: list[str] = []
        for length in (2, 4, 6, 8, 10):
            if len(code) < length:
                continue
            desc = all_desc.get(code[:length])
            if not desc:
                continue
            desc = desc.strip()
            # Названия разделов — это абзац канцелярита на 300+ знаков, одинаковый
            # для всех кодов раздела. Целиком он забивает вектор и снижает
            # различимость, поэтому каждый уровень подрезаем.
            if len(desc) > LEVEL_TEXT_LIMIT:
                desc = desc[:LEVEL_TEXT_LIMIT].rsplit(" ", 1)[0] + "…"
            # не повторяем одно и то же название на соседних уровнях
            if parts and desc.lower() == parts[-1].lower():
                continue
            parts.append(desc)

        own = (by_code[code].get("description") or "").strip()
        if own and (not parts or own.lower() != parts[-1].lower()):
            parts.append(own)

        if not parts:
            parts = [(row["full_path"] or row["description"] or "").strip()]
        texts.append(" → ".join(p for p in parts if p))

    print(f"Текстов из цепочек TWS: {used_tws:,}, собрано из иерархии базы: "
          f"{len(rows) - used_tws:,}")
    return texts


def load_leaves() -> list[dict]:
    if not DB_PATH.exists():
        print(f"База данных не найдена: {DB_PATH}")
        print("Запустите: python parse_tnved.py")
        sys.exit(1)

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    # Берём все коды, не только level=4, чтобы не зависеть от качества данных
    cur.execute("""
        SELECT code, description, full_path, duty_rate, data_source
        FROM codes
        WHERE level >= 3
        ORDER BY code
    """)
    rows = [
        {
            "code": r[0],
            "description": r[1],
            "full_path": r[2],
            "duty_rate": r[3],
            "data_source": r[4],
        }
        for r in cur.fetchall()
    ]
    conn.close()
    return rows


def stage_embed() -> None:
    """Этап 1: тексты → векторы в data/tnved_vecs.npy. faiss здесь не импортируем.

    Считаем кусками по CHUNK_SIZE и сохраняем каждый кусок отдельным .npy,
    чтобы обрыв на середине не обнулял работу: повторный запуск досчитывает
    только недостающие куски. На локальном CPU это 3–6 текстов/с (часы),
    через bge-m3 на GPU — минуты.
    """
    import numpy as np

    from embedder import active_backend, expected_name, get_embedder

    backend = active_backend()
    if backend == "local" and EMBED_THREADS:
        try:
            import torch
        except ImportError:
            sys.exit("EMBEDDER_BACKEND=local: нужен sentence-transformers — "
                     "pip install -r requirements-e5.txt")

        torch.set_num_threads(EMBED_THREADS)

    print(f"Загружаем коды из {DB_PATH} ...")
    rows = load_leaves()
    print(f"Кодов для индексирования: {len(rows):,}")

    if not rows:
        print("Нет данных для индексирования.")
        sys.exit(1)

    # Префиксы (если нужны) добавляет сам бэкенд — они у моделей разные.
    texts = build_index_texts(rows)
    uniq = len(set(texts))
    print(f"Текстов для векторизации: {len(texts):,}, из них уникальных {uniq:,} "
          f"({100 * uniq / len(texts):.1f}%)")
    for row, text in zip(rows, texts):
        row["index_text"] = text

    with open(META_PATH, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    print(f"Метаданные сохранены: {META_PATH}")

    # Отпечаток входа: и сами тексты, и модель. Сверять только количество записей
    # нельзя — при смене способа склейки текста количество то же, а векторы уже
    # не те, и старый кэш подхватывался молча.
    fingerprint = {
        "embedder": expected_name(),
        "count": len(texts),
        "texts_sha256": hashlib.sha256(
            "\n".join(texts).encode("utf-8")
        ).hexdigest(),
    }
    stored = None
    if VECS_INFO_PATH.exists():
        with open(VECS_INFO_PATH, encoding="utf-8") as f:
            stored = json.load(f)

    if VECS_PATH.exists() and stored == fingerprint:
        print(f"Векторы уже посчитаны для этих текстов и модели, пропускаем этап embed.")
        print(f"Пересчитать заново — удалите {VECS_PATH}")
        return
    if stored is not None and stored != fingerprint:
        diff = [k for k in fingerprint if stored.get(k) != fingerprint[k]]
        print(f"Кэш векторов устарел (изменилось: {', '.join(diff)}) — считаем заново.")
        if PARTS_DIR.exists():
            for old in PARTS_DIR.glob("part_*.npy"):
                old.unlink()
        VECS_PATH.unlink(missing_ok=True)

    chunks = [(i, texts[i:i + CHUNK_SIZE]) for i in range(0, len(texts), CHUNK_SIZE)]

    PARTS_DIR.mkdir(parents=True, exist_ok=True)
    todo = [(i, ch) for i, ch in chunks if not _part_path(i, len(ch)).exists()]
    done = len(chunks) - len(todo)
    print(f"\nКусков по {CHUNK_SIZE}: всего {len(chunks)}, готово {done}, к расчёту {len(todo)}")

    if todo:
        print(f"Бэкенд векторизации: {backend}")
        embedder = get_embedder()
        print(f"Модель: {embedder.name}")

        for num, (offset, chunk) in enumerate(todo, 1):
            t0 = time.time()
            vecs = embedder.encode(chunk, is_query=False)
            path = _part_path(offset, len(chunk))
            np.save(path, vecs)
            dt = time.time() - t0
            left = (len(todo) - num) * dt
            print(f"  [{num}/{len(todo)}] {path.name}  {dt:.0f}с  "
                  f"({len(chunk)/max(dt, 0.01):.0f} текстов/с, "
                  f"осталось ~{left/60:.0f} мин)", flush=True)

    print("\nСклеиваем куски ...")
    parts = [np.load(_part_path(i, len(ch))) for i, ch in chunks]
    vecs = np.concatenate(parts, axis=0)
    if vecs.shape[0] != len(rows):
        print(f"Не сходится: склеено {vecs.shape[0]}, кодов {len(rows)}. Прерываю.")
        sys.exit(1)

    np.save(VECS_PATH, vecs)
    with open(VECS_INFO_PATH, "w", encoding="utf-8") as f:
        json.dump(fingerprint, f, ensure_ascii=False, indent=2)
    print(f"Векторы сохранены: {VECS_PATH}  {vecs.shape}")
    for i, ch in chunks:
        _part_path(i, len(ch)).unlink()
    PARTS_DIR.rmdir()


def stage_index() -> None:
    """Этап 2: векторы из .npy → FAISS-индекс. torch здесь не импортируем."""
    import faiss
    import numpy as np

    from embedder import expected_name

    if not VECS_PATH.exists():
        print(f"Нет файла векторов: {VECS_PATH}")
        print("Запустите: python build_index.py embed")
        sys.exit(1)

    # В паспорт — модель, которая посчитала векторы (отпечаток этапа embed),
    # а не та, что задана сейчас: `build_index.py index` с другими настройками
    # записал бы чужое имя, и сверка паспорта на старте пропустила бы индекс.
    built_by = None
    if VECS_INFO_PATH.exists():
        with open(VECS_INFO_PATH, encoding="utf-8") as f:
            built_by = json.load(f).get("embedder")
    if not built_by:
        sys.exit(f"Нет отпечатка векторов ({VECS_INFO_PATH.name}) — неизвестно, какой моделью "
                 f"они посчитаны. Пересчитайте: python build_index.py embed")
    if built_by != expected_name():
        print(f"ВНИМАНИЕ: векторы посчитаны моделью «{built_by}», а сейчас задана "
              f"«{expected_name()}». Паспорт пишу по векторам; сервис с текущими "
              f"настройками этот индекс не примет.")

    vecs = np.load(VECS_PATH)
    print(f"Строим FAISS-индекс (IndexFlatIP, dim={vecs.shape[1]}) ...")
    index = faiss.IndexFlatIP(vecs.shape[1])
    index.add(vecs)

    faiss.write_index(index, str(FAISS_PATH))
    print(f"Индекс сохранён: {FAISS_PATH}  ({index.ntotal:,} векторов)")

    # Паспорт индекса: чем собран и в какой размерности. Без него поиск другим
    # бэкендом искал бы в чужом векторном пространстве и молча врал.
    info = {
        "embedder": built_by,
        "dim": int(vecs.shape[1]),
        "count": int(index.ntotal),
    }
    with open(INFO_PATH, "w", encoding="utf-8") as f:
        json.dump(info, f, ensure_ascii=False, indent=2)
    print(f"Паспорт индекса: {INFO_PATH}  {info}")


def main() -> None:
    stage = sys.argv[1] if len(sys.argv) > 1 else None

    if stage == "embed":
        stage_embed()
        return
    if stage == "index":
        stage_index()
        return
    if stage is not None:
        print(f"Неизвестный этап: {stage}. Допустимо: embed | index")
        sys.exit(2)

    # Без аргументов — гоняем оба этапа, каждый в отдельном процессе.
    for name in ("embed", "index"):
        print(f"\n{'=' * 60}\nЭтап: {name}\n{'=' * 60}")
        # -u обязателен: при падении процесса буферизованный stdout теряется,
        # и в логе не видно, на каком куске всё оборвалось.
        result = subprocess.run(
            [sys.executable, "-u", str(Path(__file__).resolve()), name]
        )
        if result.returncode != 0:
            print(f"\nЭтап {name} упал (код {result.returncode}). Прогон остановлен.")
            sys.exit(result.returncode)

    print("\nГотово. Индекс и метаданные на месте.")


if __name__ == "__main__":
    main()
