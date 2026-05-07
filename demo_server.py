"""
Демо-сервер для превью фронта без LLM и FAISS-индекса.
Возвращает правдоподобные заглушки на тех же эндпоинтах, что api.py.

Запуск: uvicorn demo_server:app --host 127.0.0.1 --port 8000
"""

from __future__ import annotations

import asyncio
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel

BASE = Path(__file__).parent
STATIC_DIR = BASE / "static"

app = FastAPI(title="ТН ВЭД Ассистент — DEMO")

sessions: dict[str, dict] = {}
batch_jobs: dict[str, dict] = {}
chats: dict[str, dict] = {}


# ─── мок-данные ───────────────────────────────────────────────────────────────

SAMPLE_GROUP = {"code": "84", "name": "Реакторы ядерные, котлы, оборудование и механические устройства; их части"}

SAMPLE_QUESTIONS = [
    {"id": "q1", "question": "Это бытовое или промышленное использование?",
     "hint": "От этого зависит выбор товарной позиции внутри группы 84."},
    {"id": "q2", "question": "Какие массовые/мощностные характеристики?",
     "hint": "Иногда коды разделяются по граничным значениям (например, до 10 кВт / свыше)."},
    {"id": "q3", "question": "Это полностью готовое изделие или комплект для сборки?",
     "hint": "Это влияет на применимость ОПИ-2(а)."},
]

SAMPLE_RESULT = {
    "primary": {
        "code": "8422110000",
        "reasoning": "Описание соответствует посудомоечной машине бытового типа (10–14 комплектов посуды), встраиваемой. По ОПИ-1 классификация по тексту товарной позиции 8422 11.",
        "confidence": "high",
        "hierarchy": [
            {"code": "84", "description": "Реакторы ядерные, котлы, оборудование..."},
            {"code": "8422", "description": "Машины посудомоечные; оборудование..."},
            {"code": "842211", "description": "посудомоечные машины бытового типа"},
        ],
        "full_path": "Реакторы ядерные, котлы, оборудование... → Машины посудомоечные → бытовые → встраиваемые",
        "duty_rate": "5%",
    },
    "alternatives": [
        {"code": "8422190000", "why_close": "Близкая позиция на промышленные посудомойки",
         "why_rejected": "Описание явно бытовое (12 комплектов, встраиваемая)",
         "full_path": "Машины посудомоечные → промышленные", "duty_rate": "0%"},
        {"code": "8421120000", "why_close": "Иногда путают со стиральными",
         "why_rejected": "Стиральные машины — отдельная позиция 8450",
         "full_path": "Сушки центробежные для белья", "duty_rate": "10%"},
    ],
    "gri_explained": [
        {"code": "1", "text": "Классификация определяется в соответствии с текстами товарных позиций..."},
    ],
    "checks_required": [
        "Уточните соответствие количества комплектов диапазону позиции 8422 11",
        "Если в комплекте идут аксессуары — проверьте применимость ОПИ-3(б)",
    ],
    "candidates": [],
}


def _delay():
    """Имитируем задержку LLM."""
    return asyncio.sleep(0.8)


# ─── эндпоинты, повторяющие api.py ────────────────────────────────────────────

class StartRequest(BaseModel):
    mode: str = "simple"
    description: str | None = None
    fields: dict | None = None


class FinalizeRequest(BaseModel):
    session_id: str
    answers: list = []


@app.get("/api/models")
async def models():
    """Демо-сервер делает вид, что поддерживает оба провайдера. UI покажет переключение."""
    return {
        "default": {"provider": "openai", "model": "demo-model"},
        "providers": {
            "openai": {
                "label": "OpenAI",
                "models": ["gpt-4o-mini", "gpt-4o", "gpt-4-turbo"],
                "needs_base_url": True,
            },
            "anthropic": {
                "label": "Anthropic",
                "models": ["claude-haiku-4-5", "claude-sonnet-4-6", "claude-opus-4-7"],
                "needs_base_url": False,
            },
        },
    }


@app.post("/api/classify/start")
async def classify_start(req: StartRequest):
    await _delay()
    desc = (req.description or "").strip()
    if not desc and req.fields:
        desc = "; ".join(f"{k}: {v}" for k, v in req.fields.items() if v)
    if not desc:
        raise HTTPException(400, "Пустое описание")
    sid = uuid.uuid4().hex[:12]
    sessions[sid] = {
        "id": sid, "description": desc,
        "created_at": datetime.utcnow().isoformat(),
        "triage": {"group_code": SAMPLE_GROUP["code"], "group_name": SAMPLE_GROUP["name"],
                   "completeness": "low", "missing_aspects": ["назначение", "тех. параметры"],
                   "questions": SAMPLE_QUESTIONS},
    }
    return {
        "session_id": sid,
        "description": desc,
        "group": SAMPLE_GROUP,
        "completeness": "low",
        "missing_aspects": ["назначение", "технические характеристики"],
        "questions": SAMPLE_QUESTIONS,
    }


@app.post("/api/classify/finalize")
async def classify_finalize(req: FinalizeRequest):
    await _delay()
    session = sessions.get(req.session_id)
    if not session:
        raise HTTPException(404, "Сессия не найдена")
    session["result"] = SAMPLE_RESULT
    return {
        "session_id": req.session_id,
        "description": session["description"],
        "group": SAMPLE_GROUP,
        "result": SAMPLE_RESULT,
    }


@app.get("/api/classify/{session_id}")
async def classify_get(session_id: str):
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(404, "Сессия не найдена")
    return session


# ─── batch ────────────────────────────────────────────────────────────────────

@app.post("/api/classify/batch")
async def classify_batch_start(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Ожидается .xlsx")
    content = await file.read()
    try:
        wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(400, f"Не смог прочитать xlsx: {e}")
    descriptions = []
    for row in rows[1:] if rows else []:
        if row and row[0]:
            descriptions.append(str(row[0]).strip())
    if not descriptions and rows:
        for row in rows:
            if row and row[0]:
                descriptions.append(str(row[0]).strip())
    if not descriptions:
        raise HTTPException(400, "Нет описаний")

    job_id = uuid.uuid4().hex[:12]
    batch_jobs[job_id] = {
        "id": job_id, "filename": file.filename, "total": len(descriptions),
        "processed": 0, "status": "running", "errors": [],
        "started_at": datetime.utcnow().isoformat(), "finished_at": None,
        "descriptions": descriptions,
    }

    async def run():
        for i in range(len(descriptions)):
            await asyncio.sleep(0.5)
            batch_jobs[job_id]["processed"] = i + 1
        batch_jobs[job_id]["status"] = "done"
        batch_jobs[job_id]["finished_at"] = datetime.utcnow().isoformat()

    asyncio.create_task(run())
    return {"job_id": job_id, "total": len(descriptions), "status": "running"}


@app.get("/api/classify/batch/{job_id}")
async def classify_batch_status(job_id: str):
    job = batch_jobs.get(job_id)
    if not job:
        raise HTTPException(404, "Job не найден")
    return {
        "job_id": job["id"], "filename": job["filename"], "total": job["total"],
        "processed": job["processed"], "status": job["status"],
        "errors_count": len(job["errors"]),
        "started_at": job["started_at"], "finished_at": job["finished_at"],
    }


@app.get("/api/classify/batch/{job_id}/download")
async def classify_batch_download(job_id: str):
    job = batch_jobs.get(job_id)
    if not job:
        raise HTTPException(404, "Job не найден")
    if job["status"] != "done":
        raise HTTPException(409, "Job ещё не завершён")
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат"
    ws.append(["№", "Описание", "Код ТН ВЭД", "Наименование", "Пошлина",
               "Уверенность", "Группа", "Альт. 1", "Альт. 1 пошлина",
               "Альт. 2", "Альт. 2 пошлина", "Ошибка"])
    for i, desc in enumerate(job["descriptions"], 1):
        ws.append([i, desc, "8422110000", "посудомоечные машины бытовые встраиваемые",
                   "5%", "high", f'{SAMPLE_GROUP["code"]} {SAMPLE_GROUP["name"]}',
                   "8422190000", "0%", "8421120000", "10%", ""])
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"tnved_batch_{job_id}.xlsx"
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ─── chat ─────────────────────────────────────────────────────────────────────

class ChatStartRequest(BaseModel):
    initial_description: str = ""


class ChatMessageRequest(BaseModel):
    text: str


def _bot_question_msg(turn: int) -> str:
    bank = [
        f"Похоже на группу **84** — оборудование и механические устройства.\nУточню:\n- Бытовое или промышленное использование?\n- Технические параметры (мощность, объём)?",
        f"Спасибо. Ещё вопрос:\n- Это полностью готовое изделие или комплект?\n- Есть ли встраиваемое исполнение?",
        f"Почти готово, последний вопрос:\n- Сколько комплектов посуды загружается одновременно?",
    ]
    return bank[min(turn, len(bank) - 1)]


def _bot_final_msg() -> str:
    return (
        "**Код:** `8422110000`\n"
        "_Машины посудомоечные → бытовые → встраиваемые_\n"
        "**Пошлина:** 5%  **Уверенность:** high\n"
        "\n"
        "Описание соответствует посудомоечной машине бытового типа (10–14 комплектов), встраиваемой. По ОПИ-1 — позиция 8422 11.\n"
        "\n"
        "**Альтернативы:**\n"
        "- `8422190000` (0%) — близкая позиция на промышленные посудомойки\n"
        "- `8421120000` (10%) — иногда путают со стиральными\n"
        "\n"
        "**Стоит проверить вручную:**\n"
        "- Соответствие количества комплектов диапазону позиции 8422 11\n"
        "- Наличие аксессуаров — может сработать ОПИ-3(б)"
    )


@app.post("/api/chat/start")
async def chat_start(req: ChatStartRequest):
    cid = uuid.uuid4().hex[:12]
    chats[cid] = {"id": cid, "messages": [], "phase": "gathering",
                  "user_turns": 0, "description": ""}
    if req.initial_description.strip():
        return await _chat_handle(cid, req.initial_description.strip())
    chats[cid]["messages"].append({
        "role": "assistant",
        "content": "Опишите товар — состав, назначение, форму, технические параметры. Я задам уточняющие вопросы и подберу код ТН ВЭД.",
    })
    return {"chat_id": cid, "phase": "gathering",
            "messages": chats[cid]["messages"], "result": None}


async def _chat_handle(cid: str, text: str):
    await _delay()
    chat = chats[cid]
    chat["messages"].append({"role": "user", "content": text})
    chat["user_turns"] += 1

    # после 3-го хода — финализация
    if chat["user_turns"] >= 3:
        chat["messages"].append({"role": "assistant", "content": _bot_final_msg()})
        chat["phase"] = "finalized"
    else:
        chat["messages"].append({
            "role": "assistant",
            "content": _bot_question_msg(chat["user_turns"] - 1),
        })
    return {"chat_id": cid, "phase": chat["phase"],
            "messages": chat["messages"], "result": None}


@app.post("/api/chat/{chat_id}/message")
async def chat_message(chat_id: str, req: ChatMessageRequest):
    if chat_id not in chats:
        raise HTTPException(404, "Чат не найден")
    if chats[chat_id]["phase"] == "finalized":
        raise HTTPException(409, "Чат уже завершён")
    return await _chat_handle(chat_id, req.text)


@app.get("/api/chat/{chat_id}")
async def chat_get(chat_id: str):
    if chat_id not in chats:
        raise HTTPException(404, "Чат не найден")
    return chats[chat_id]


# ─── статика ──────────────────────────────────────────────────────────────────

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/")
async def index():
    return FileResponse(STATIC_DIR / "index.html")
